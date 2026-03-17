#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génération du document de spécification e-SONABE pour la Direction de SONABEL
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# --- Couleurs SONABEL ---
RED_SONABEL   = RGBColor(0xD3, 0x2F, 0x2F)   # #D32F2F
GREEN_SONABEL = RGBColor(0x2E, 0x7D, 0x32)   # #2E7D32
YELLOW        = RGBColor(0xFF, 0xD6, 0x00)   # #FFD600
DARK_GREY     = RGBColor(0x21, 0x21, 0x21)
LIGHT_GREY    = RGBColor(0xF5, 0xF5, 0xF5)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)

def set_cell_bg(cell, color: RGBColor):
    """Définit la couleur de fond d'une cellule."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def add_horizontal_rule(doc):
    """Ajoute un filet horizontal."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E7D32')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)

def add_heading(doc, text, level=1, color=GREEN_SONABEL):
    """Ajoute un titre formaté."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(20)
    elif level == 2:
        run.font.size = Pt(15)
    else:
        run.font.size = Pt(12)
    run.font.color.rgb = color
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_body(doc, text, bold=False, italic=False, size=11, color=DARK_GREY):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_GREY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    return p

def add_feature_table(doc, features):
    """Tableau à deux colonnes : icône-titre | description."""
    table = doc.add_table(rows=len(features), cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (title, desc) in enumerate(features):
        row = table.rows[i]
        # Col 1 – titre
        c1 = row.cells[0]
        c1.width = Cm(5.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(title)
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = GREEN_SONABEL
        if i % 2 == 0:
            set_cell_bg(c1, RGBColor(0xE8, 0xF5, 0xE9))
        else:
            set_cell_bg(c1, RGBColor(0xF1, 0xF8, 0xE9))
        # Col 2 – description
        c2 = row.cells[1]
        c2.width = Cm(10.5)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(desc)
        r2.font.size = Pt(11)
        r2.font.color.rgb = DARK_GREY
        if i % 2 == 0:
            set_cell_bg(c2, RGBColor(0xFA, 0xFF, 0xFA))
        else:
            set_cell_bg(c2, RGBColor(0xF9, 0xFD, 0xF9))
    doc.add_paragraph()
    return table

def add_value_prop_table(doc, props):
    """Tableau de la proposition de valeur."""
    table = doc.add_table(rows=1 + len(props), cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Problème Actuel", "Solution e-SONABE", "Bénéfice Mesurable"]
    hdr_row = table.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        set_cell_bg(cell, GREEN_SONABEL)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = WHITE
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (prob, sol, ben) in enumerate(props):
        row = table.rows[i + 1]
        bg = RGBColor(0xFA, 0xFA, 0xFA) if i % 2 == 0 else WHITE
        for j, text in enumerate([prob, sol, ben]):
            cell = row.cells[j]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.color.rgb = DARK_GREY
    doc.add_paragraph()
    return table

# ============================================================
# Création du document
# ============================================================
doc = Document()

# --- Marges ---
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ============================================================
# PAGE DE COUVERTURE
# ============================================================

# Bandeau vert supérieur simulé via tableau
cover_top = doc.add_table(rows=1, cols=1)
cover_top.alignment = WD_TABLE_ALIGNMENT.CENTER
ct_cell = cover_top.rows[0].cells[0]
set_cell_bg(ct_cell, GREEN_SONABEL)
ct_cell.height = Cm(0.5)
ct_p = ct_cell.paragraphs[0]
ct_p.add_run(" ").font.size = Pt(6)
doc.add_paragraph()

# Logo / Titre principal
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("e-SONABE")
title_run.bold = True
title_run.font.size = Pt(48)
title_run.font.color.rgb = GREEN_SONABEL

subtitle_p = doc.add_paragraph()
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = subtitle_p.add_run("Votre Espace Client Électricité")
sub_run.font.size = Pt(18)
sub_run.font.color.rgb = RED_SONABEL
sub_run.bold = True

doc.add_paragraph()

# Filet coloré
add_horizontal_rule(doc)

doc.add_paragraph()

# Titre du document
doc_title = doc.add_paragraph()
doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
dt_run = doc_title.add_run("DOCUMENT DE SPÉCIFICATION\nPRODUIT & PROPOSITION DE VALEUR")
dt_run.bold = True
dt_run.font.size = Pt(16)
dt_run.font.color.rgb = DARK_GREY

doc.add_paragraph()
doc.add_paragraph()

# Destinataire
dest_table = doc.add_table(rows=3, cols=2)
dest_table.alignment = WD_TABLE_ALIGNMENT.CENTER
dest_data = [
    ("Destinataire :", "Direction Générale de SONABEL"),
    ("Objet :", "Présentation de la solution e-SONABE"),
    ("Date :", datetime.date.today().strftime("%d %B %Y")),
]
for i, (lbl, val) in enumerate(dest_data):
    row = dest_table.rows[i]
    c1 = row.cells[0]
    c1.width = Cm(4)
    set_cell_bg(c1, RGBColor(0xE8, 0xF5, 0xE9))
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(lbl)
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = GREEN_SONABEL

    c2 = row.cells[1]
    c2.width = Cm(9)
    p2 = c2.paragraphs[0]
    r2 = p2.add_run(val)
    r2.font.size = Pt(11)
    r2.font.color.rgb = DARK_GREY

doc.add_paragraph()
doc.add_paragraph()

# Bandeau rouge inférieur
cover_bot = doc.add_table(rows=1, cols=1)
cover_bot.alignment = WD_TABLE_ALIGNMENT.CENTER
cb_cell = cover_bot.rows[0].cells[0]
set_cell_bg(cb_cell, RED_SONABEL)
cb_p = cb_cell.paragraphs[0]
cb_run = cb_p.add_run("CONFIDENTIEL – USAGE INTERNE SONABEL")
cb_run.bold = True
cb_run.font.size = Pt(9)
cb_run.font.color.rgb = WHITE
cb_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Saut de page
doc.add_page_break()

# ============================================================
# 1. RÉSUMÉ EXÉCUTIF
# ============================================================
add_heading(doc, "1. Résumé Exécutif", level=1)
add_horizontal_rule(doc)
add_body(doc,
    "Dans un contexte où la digitalisation des services publics représente un enjeu stratégique "
    "majeur, SONABEL a l'opportunité de franchir un cap décisif en matière de relation client. "
    "e-SONABE est une plateforme web moderne, développée sur mesure, qui permet à chaque client "
    "de SONABEL de gérer en toute autonomie ses contrats d'électricité, ses compteurs, ses "
    "factures et sa consommation — depuis n'importe quel appareil connecté, à tout moment.")

add_body(doc,
    "Pensée pour le marché burkinabè et ouest-africain, la solution intègre nativement les "
    "compteurs à prépaiement CashPower, une interface bilingue français/anglais, et un module "
    "d'analyse prédictive de la consommation propulsé par intelligence artificielle.")

add_body(doc,
    "Ce document présente les spécifications fonctionnelles, les avantages compétitifs et la "
    "proposition de valeur d'e-SONABE à l'intention de la Direction Générale de SONABEL.",
    italic=True, color=RGBColor(0x55, 0x55, 0x55))

# ============================================================
# 2. CONTEXTE & PROBLÉMATIQUE
# ============================================================
doc.add_page_break()
add_heading(doc, "2. Contexte & Problématique", level=1)
add_horizontal_rule(doc)
add_body(doc,
    "Aujourd'hui, les clients de SONABEL font face à plusieurs obstacles dans la gestion de "
    "leurs abonnements électriques :")

problems = [
    "Obligation de se déplacer en agence pour toute demande (consultation de facture, paiement, déclaration de panne).",
    "Absence de visibilité en temps réel sur la consommation et les dépenses d'électricité.",
    "Difficulté à gérer plusieurs points de livraison (contrats multiples) depuis un seul endroit.",
    "Manque d'outils de prévision permettant d'anticiper les pics de consommation et les coûts.",
    "Processus de rechargement CashPower peu dématérialisé, exposant les clients à des risques de coupure.",
    "Absence d'historique numérique accessible pour les litiges ou les vérifications.",
]
for pb in problems:
    add_bullet(doc, pb)

add_body(doc,
    "\nCes lacunes engendrent une insatisfaction client croissante, une charge opérationnelle "
    "élevée pour les agences SONABEL et une image d'entreprise en décalage avec les standards "
    "modernes du secteur de l'énergie.", italic=True)

# ============================================================
# 3. DESCRIPTION DE LA SOLUTION
# ============================================================
doc.add_page_break()
add_heading(doc, "3. Description de la Solution e-SONABE", level=1)
add_horizontal_rule(doc)
add_body(doc,
    "e-SONABE est un portail client web full-stack développé avec le framework Laravel (PHP 8.2+), "
    "une interface Tailwind CSS réactive et des composants Alpine.js. La solution s'appuie sur "
    "une architecture robuste, sécurisée et facilement déployable sur l'infrastructure de SONABEL.")

add_heading(doc, "3.1 Modules Fonctionnels", level=2)

features = [
    ("🔐  Espace Client Sécurisé",
     "Inscription, connexion avec vérification par e-mail, gestion de profil, "
     "changement de mot de passe et suppression de compte sécurisée."),
    ("📋  Gestion des Contrats",
     "Création, consultation, modification et résiliation des contrats d'électricité. "
     "Suivi du statut (actif / suspendu / résilié) pour chaque point de livraison."),
    ("⚡  Gestion des Compteurs",
     "Enregistrement des compteurs normaux (postpayés) et CashPower (prépayés) avec "
     "classes A, B1, B2, C1, C2. Génération automatique des numéros de série."),
    ("🧾  Facturation & Paiements",
     "Consultation des factures avec détail (consommation kWh, montant, taxes, TVA), "
     "filtrage par statut (en attente / payée / en retard) et paiement en ligne."),
    ("🔋  Recharge CashPower",
     "Achat de crédit électrique prépayé de 500 à 500 000 FCFA. "
     "Génération de tokens sécurisés et historique complet des transactions."),
    ("📊  Tableau de Bord KPI",
     "Vue synthétique : nombre de contrats actifs, compteurs enregistrés, "
     "factures en attente et montant total dû — en un seul coup d'œil."),
    ("📈  Analyse de Consommation",
     "Visualisation graphique de la consommation sur 30 jours, "
     "statistiques avancées (moyenne journalière, pic de consommation, tendance)."),
    ("🤖  IA & Prévision",
     "Module de prévision par lissage exponentiel avec détection de saisonnalité hebdomadaire. "
     "Projections à 7 et 30 jours avec indice de confiance et recommandations contextuelles."),
    ("🌍  Interface Bilingue",
     "Basculement instantané Français / Anglais, adapté aux profils "
     "d'utilisateurs burkinabè et de la sous-région ouest-africaine."),
    ("📱  Design Responsive",
     "Interface optimisée pour ordinateur, tablette et smartphone — "
     "accessible sans installation depuis n'importe quel navigateur."),
]
add_feature_table(doc, features)

# ============================================================
# 4. AVANTAGES DU PRODUIT
# ============================================================
doc.add_page_break()
add_heading(doc, "4. Avantages du Produit", level=1)
add_horizontal_rule(doc)

add_heading(doc, "4.1 Pour les Clients de SONABEL", level=2)
client_advantages = [
    "Autonomie totale 24h/24, 7j/7 — plus besoin de se déplacer en agence pour les démarches courantes.",
    "Visibilité complète sur leur consommation et leurs dépenses en temps réel.",
    "Anticipation des coupures grâce aux alertes de solde bas sur les compteurs CashPower (seuil < 20 kWh).",
    "Maîtrise budgétaire grâce au coût mensuel projeté par l'IA.",
    "Gestion centralisée de plusieurs contrats et compteurs depuis un seul compte.",
    "Historique numérique complet des factures et des transactions — consultable à tout moment.",
    "Expérience utilisateur fluide et moderne, dans leur langue (français ou anglais).",
]
for a in client_advantages:
    add_bullet(doc, a)

add_heading(doc, "4.2 Pour SONABEL (Opérateur)", level=2)
operator_advantages = [
    "Réduction significative du flux en agence et des coûts de traitement manuel.",
    "Diminution des impayés grâce aux rappels automatiques et à la visibilité des factures en retard.",
    "Données de consommation agrégées pour une meilleure planification de la production et du réseau.",
    "Image de marque modernisée, alignée sur les standards internationaux du secteur de l'énergie.",
    "Plateforme évolutive, prête pour l'intégration future de nouveaux services (facturation automatique, API partenaires, paiement mobile Money).",
    "Conformité aux bonnes pratiques de sécurité (authentification, vérification e-mail, contrôle d'accès par politique).",
    "Base technologique Laravel 12 — framework éprouvé, large écosystème, maintenance facilitée.",
]
for a in operator_advantages:
    add_bullet(doc, a)

add_heading(doc, "4.3 Avantages Technologiques", level=2)
tech_advantages = [
    "Architecture MVC claire avec séparation stricte des responsabilités — maintenable et extensible.",
    "Tests automatisés intégrés (PHPUnit) pour garantir la qualité et faciliter les évolutions.",
    "Environnement Docker (Laravel Sail) pour un déploiement rapide et reproductible.",
    "Système de migration de base de données versionné — traçabilité complète des évolutions du schéma.",
    "Interface API RESTful prête pour l'intégration avec des systèmes tiers (ERP, CRM, paiement mobile).",
    "Localisation i18n native — ajout simple de nouvelles langues sans refonte du code.",
    "Optimisation frontend avec Vite + Tailwind CSS — temps de chargement réduit, expérience fluide même avec un réseau limité.",
]
for a in tech_advantages:
    add_bullet(doc, a)

# ============================================================
# 5. PROPOSITION DE VALEUR
# ============================================================
doc.add_page_break()
add_heading(doc, "5. Proposition de Valeur", level=1)
add_horizontal_rule(doc)

add_body(doc,
    "e-SONABE transforme la relation entre SONABEL et ses clients en remplaçant les interactions "
    "physiques coûteuses par une expérience digitale fluide, intelligente et disponible en permanence. "
    "La proposition de valeur repose sur trois piliers :")

doc.add_paragraph()

# Trois piliers en tableau
pillars = doc.add_table(rows=1, cols=3)
pillars.alignment = WD_TABLE_ALIGNMENT.CENTER
pillar_data = [
    ("SIMPLICITÉ", GREEN_SONABEL,
     "Un portail unique pour tout gérer : contrats, compteurs, factures, recharges — accessible en 3 clics."),
    ("INTELLIGENCE", RED_SONABEL,
     "L'IA anticipe votre consommation, vous alerte avant une coupure et optimise votre budget énergie."),
    ("CONFIANCE", RGBColor(0x15, 0x65, 0xC0),
     "Données sécurisées, historique fiable, transactions traçables — une relation client transparente."),
]
for j, (title, color, desc) in enumerate(pillar_data):
    cell = pillars.rows[0].cells[j]
    set_cell_bg(cell, color)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"\n{title}\n\n").font.color.rgb = WHITE
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(13)
    r2 = p.add_run(desc)
    r2.font.size = Pt(10)
    r2.font.color.rgb = WHITE
    p.add_run("\n")

doc.add_paragraph()

add_heading(doc, "5.1 Tableau Valeur : Problème → Solution → Bénéfice", level=2)
value_props = [
    ("Déplacements fréquents en agence",
     "Portail 100% en ligne, disponible 24h/24",
     "Gain de temps estimé à 2-4h par démarche client"),
    ("Factures papier perdues ou illisibles",
     "Espace numérique avec historique complet et filtres",
     "Réduction des litiges factures de ~60%"),
    ("Coupures CashPower imprévues",
     "Alertes solde bas + recharge instantanée en ligne",
     "Continuité de service, satisfaction client accrue"),
    ("Consommation opaque, budget non maîtrisé",
     "Graphiques, statistiques et prévisions IA à 30 jours",
     "Économies d'énergie potentielles de 10 à 20%"),
    ("Gestion manuelle multi-contrats complexe",
     "Tableau de bord centralisé multi-contrats/compteurs",
     "Productivité des gestionnaires d'immeubles améliorée"),
    ("Charge en agence et coûts opérationnels élevés",
     "Digitalisation des processus clés",
     "Réduction des coûts opérationnels SONABEL de ~30%"),
]
add_value_prop_table(doc, value_props)

# ============================================================
# 6. SPÉCIFICATIONS TECHNIQUES
# ============================================================
doc.add_page_break()
add_heading(doc, "6. Spécifications Techniques", level=1)
add_horizontal_rule(doc)

tech_table = doc.add_table(rows=8, cols=2)
tech_table.style = 'Table Grid'
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
tech_specs = [
    ("Framework Backend",   "Laravel 12 (PHP 8.2+)"),
    ("Interface Frontend",  "Blade Templates, Tailwind CSS 3.1, Alpine.js 3.4"),
    ("Base de Données",     "SQLite (déploiement rapide) / MySQL / PostgreSQL"),
    ("Outil de Build",      "Vite 7.0 + PostCSS + Autoprefixer"),
    ("Authentification",    "Laravel Breeze — vérification e-mail, sessions sécurisées"),
    ("Tests",               "PHPUnit + Mockery — couverture des fonctionnalités critiques"),
    ("Déploiement",         "Laravel Sail (Docker) — compatible serveurs Linux"),
    ("Devise",              "Franc CFA (XOF) — précision décimale sur les montants financiers"),
]
for i, (label, value) in enumerate(tech_specs):
    row = tech_table.rows[i]
    c1 = row.cells[0]
    c1.width = Cm(5.5)
    set_cell_bg(c1, RGBColor(0xE8, 0xF5, 0xE9))
    r1 = c1.paragraphs[0].add_run(label)
    r1.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = GREEN_SONABEL
    c2 = row.cells[1]
    c2.width = Cm(10.5)
    r2 = c2.paragraphs[0].add_run(value)
    r2.font.size = Pt(10)
    r2.font.color.rgb = DARK_GREY

doc.add_paragraph()

# ============================================================
# 7. FEUILLE DE ROUTE & ÉVOLUTIONS ENVISAGÉES
# ============================================================
add_heading(doc, "7. Feuille de Route & Évolutions Envisagées", level=1)
add_horizontal_rule(doc)

roadmap = [
    ("Phase 1 — Déploiement Initial (J+0 à J+90)",
     ["Mise en production sur l'infrastructure SONABEL",
      "Migration des données clients existants",
      "Formation des équipes support SONABEL",
      "Campagne de communication client"]),
    ("Phase 2 — Intégration Paiement Mobile (J+90 à J+180)",
     ["Intégration Orange Money / Moov Money / Coris Money",
      "Paiement de factures directement depuis le portail",
      "Notifications SMS pour les échéances de paiement"]),
    ("Phase 3 — API Partenaires & App Mobile (J+180 à J+365)",
     ["API RESTful publique pour partenaires (banques, assurances)",
      "Application mobile iOS/Android native",
      "Tableau de bord analytique pour la Direction de SONABEL"]),
]

for phase_title, phase_items in roadmap:
    add_body(doc, phase_title, bold=True, color=GREEN_SONABEL)
    for item in phase_items:
        add_bullet(doc, item)
    doc.add_paragraph()

# ============================================================
# 8. CONCLUSION & APPEL À L'ACTION
# ============================================================
doc.add_page_break()
add_heading(doc, "8. Conclusion & Appel à l'Action", level=1)
add_horizontal_rule(doc)

add_body(doc,
    "e-SONABE représente une opportunité unique pour SONABEL de se positionner comme un acteur "
    "de référence de la transition numérique dans le secteur de l'énergie en Afrique de l'Ouest. "
    "La solution est fonctionnelle, testée et prête à être déployée.")

add_body(doc,
    "Au-delà d'un simple outil digital, e-SONABE incarne une nouvelle promesse client : "
    "celle d'une énergie gérée simplement, intelligemment, et en toute transparence.")

doc.add_paragraph()

# Encadré vert de conclusion
cta_table = doc.add_table(rows=1, cols=1)
cta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cta_cell = cta_table.rows[0].cells[0]
set_cell_bg(cta_cell, RGBColor(0xE8, 0xF5, 0xE9))
cta_p = cta_cell.paragraphs[0]
cta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cta_r1 = cta_p.add_run("Nous invitons la Direction Générale de SONABEL\n")
cta_r1.bold = True
cta_r1.font.size = Pt(13)
cta_r1.font.color.rgb = GREEN_SONABEL
cta_r2 = cta_p.add_run(
    "à nous accorder une démonstration en live de la plateforme e-SONABE\n"
    "afin d'évaluer ensemble la valeur que cette solution apportera\n"
    "à vos clients et à vos équipes opérationnelles."
)
cta_r2.font.size = Pt(12)
cta_r2.font.color.rgb = DARK_GREY
cta_p.add_run("\n")

doc.add_paragraph()

add_body(doc,
    "Pour toute information complémentaire ou pour organiser une démonstration,\n"
    "veuillez contacter l'équipe e-SONABE.",
    italic=True, color=RGBColor(0x55, 0x55, 0x55))

# ============================================================
# PIED DE PAGE
# ============================================================
for section in doc.sections:
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"e-SONABE — Document Confidentiel — {datetime.date.today().strftime('%d/%m/%Y')} — "
        "Propriété de SONABEL. Reproduction interdite sans autorisation."
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    footer_run.italic = True

# ============================================================
# SAUVEGARDE
# ============================================================
output_path = "/home/user/Cursor_folder/eSONABE_Specification_SONABEL.docx"
doc.save(output_path)
print(f"Document généré : {output_path}")
